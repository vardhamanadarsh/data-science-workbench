import io
import json

import numpy as np
import polars as pl
import pytest

from core.data_loader import load_dataset
from core.data_quality import assess_data_health
from core.eda import run_eda
from core.insights import build_context, key_findings
from core.model_io import (
    ModelIOError, can_save, inspect_untrusted, load_model_bytes, predict_from_frame, save_model_bytes,
)
from core.model_training import TrainConfig, best_result, train_models
from core.reports import TEMPLATES, build_excel, build_html, build_json, latex_escape, latex_tables


@pytest.fixture
def state(tmp_path):
    rng = np.random.default_rng(0)
    n = 300
    x1, x2 = rng.normal(size=n), rng.normal(size=n)
    y = (x1 + 0.5 * x2 + rng.normal(scale=0.5, size=n) > 0).astype(int)
    path = tmp_path / "demo.csv"
    pl.DataFrame({"x1": x1, "x2": x2, "y": y}).write_csv(path)
    ds = load_dataset(path)
    cfg = TrainConfig(target="y", features=["x1", "x2"], task="classification", cv_folds=3)
    run = train_models(ds.df, cfg, ["Logistic Regression"])
    return {"dataset": ds, "health": {"report": assess_data_health(ds.df, target="y")},
            "eda": {"result": run_eda(ds.df)}, "ml_run": run}


def test_key_findings_are_computed(state) -> None:
    text = " ".join(key_findings(build_context(state)))
    assert "data health score" in text and "Logistic Regression" in text and "Baseline" in text


def test_html_templates(state) -> None:
    ctx = build_context(state)
    for template in TEMPLATES:
        html = build_html(ctx, template)
        assert "Logistic Regression" in html and ctx.title in html


def test_html_escapes_user_text(state) -> None:
    ctx = build_context(state, "A <b>title</b>", {"general": "<script>alert(1)</script>"})
    html = build_html(ctx, "Executive summary")
    assert "<script>" not in html and "&lt;script&gt;" in html and "<b>title</b>" not in html


def test_empty_context() -> None:
    ctx = build_context({})
    assert ctx.is_empty and key_findings(ctx) == []
    assert "Nothing to report yet" in build_html(ctx)


def test_json_is_valid(state) -> None:
    text = build_json(build_context(state))
    data = json.loads(text)
    assert data["key_findings"] and data["models"][0]["results"] and "NaN" not in text


def test_latex(state) -> None:
    assert latex_escape("a_b & 5%") == r"a\_b \& 5\%"
    tables = latex_tables(build_context(state))
    assert any("Model comparison" in k for k in tables)
    table = next(iter(tables.values()))
    assert "\\toprule" in table and "\\begin{table}" in table


def test_excel(state) -> None:
    pytest.importorskip("xlsxwriter")
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(build_excel(build_context(state))))
    assert "Summary" in wb.sheetnames and "Models" in wb.sheetnames and "Dataset" in wb.sheetnames


def test_model_roundtrip(state) -> None:
    pytest.importorskip("skops")
    run = state["ml_run"]
    res = best_result(run)
    assert can_save(res)
    data = save_model_bytes(run, res)
    bundle = load_model_bytes(data, inspect_untrusted(data))
    ds = state["dataset"]
    out = predict_from_frame(bundle, ds.df)
    expected = [run.classes[int(i)] for i in res.pipeline.predict(ds.df.select(["x1", "x2"]).to_pandas())]
    assert out.height == 300 and out["prediction"].to_list() == expected
    assert "probability_1" in out.columns
    with pytest.raises(ModelIOError):
        predict_from_frame(bundle, pl.DataFrame({"x1": [1.0]}))


def test_garbage_model_file_is_rejected() -> None:
    with pytest.raises(ModelIOError):
        inspect_untrusted(b"garbage")